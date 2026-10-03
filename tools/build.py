#!/usr/bin/env python3
"""Single source of truth for the cmux Academy.

Run:  python3 tools/build.py          (writes assets/data.js)
      <venv>/bin/python tools/build.py  (also writes cmux-academy.xlsx, needs openpyxl)

Edit the data below, rebuild, and the site and spreadsheet stay in sync.
Sources: cmux.com/docs (keyboard-shortcuts, concepts, configuration, notifications,
api, browser-automation, custom-commands, workspace-groups, ssh) and `cmux --help`
for cmux 0.64.x.
"""
import json
import pathlib
import sys

ROOT = pathlib.Path(__file__).resolve().parent.parent

MODULES = [
    dict(id="m1", n=1, title="Orientation", minutes=20,
         blurb="Learn the vocabulary (window, workspace, pane, surface, panel), find yourself in the tree, and get around with the sidebar and command palette."),
    dict(id="m2", n=2, title="Layout and navigation", minutes=35,
         blurb="Splits, tabs, workspaces and groups. Build any layout from the keyboard, then again from the CLI."),
    dict(id="m3", n=3, title="Working with agents", minutes=45,
         blurb="Why cmux exists: keep several Claude Code sessions in view, get told when one needs you, and review what they changed."),
    dict(id="m4", n=4, title="The built-in browser", minutes=30,
         blurb="Open docs and localhost next to your terminal, then script the browser from the command line."),
    dict(id="m5", n=5, title="Automate and customise", minutes=60,
         blurb="Drive terminals from scripts, configure Ghostty and cmux.json, rebind keys, and define your own workspace layouts."),
]

# level: 1 = beginner, 2 = intermediate, 3 = stretch
TASKS = [
    # ---- Module 1
    dict(id="m1-1", m="m1", level=1, title="Map the hierarchy",
         steps=["Run `cmux tree`.", "Label each line as a window, workspace, pane or surface."],
         check="You can point at your own workspace and surface in the output."),
    dict(id="m1-2", m="m1", level=1, title="Find yourself",
         steps=["Run `cmux identify`.", "Run `echo $CMUX_WORKSPACE_ID $CMUX_SURFACE_ID`."],
         check="You know which workspace and surface this shell lives in, and where those ids come from."),
    dict(id="m1-3", m="m1", level=1, title="Open both sidebars",
         steps=["Press ⌘B to toggle the left sidebar (your workspaces).", "Press ⌘⌥B to toggle the right sidebar.", "Look at the modes the right sidebar offers (files, find, sessions, feed, dock...)."],
         check="You can show and hide each sidebar without the mouse."),
    dict(id="m1-4", m="m1", level=1, title="Use the command palette",
         steps=["Press ⌘⇧P.", "Search for \"reload\" and for \"split\" without running anything.", "Press ⌘P (go to workspace) and jump to another workspace."],
         check="You found a command in the palette and know how it differs from ⌘P."),
    dict(id="m1-5", m="m1", level=1, title="Locate your settings",
         steps=["Run `cmux settings path` and `cmux config paths`.", "Press ⌘, to open Settings.", "Run `cmux welcome` if you want the built-in tour."],
         check="You can say where cmux.json and the Ghostty config live."),
    # ---- Module 2
    dict(id="m2-1", m="m2", level=1, title="Split and move focus",
         steps=["⌘D splits right. ⌘⇧D splits down.", "Move focus with ⌥⌘ and the arrow keys.", "Close a pane's tab with ⌘W."],
         check="You have a three-pane layout and can reach every pane by keyboard."),
    dict(id="m2-2", m="m2", level=1, title="Zoom and equalise",
         steps=["Press ⌘⇧↩ to zoom the focused pane, and again to restore it.", "Resize with ⌃⇧H / J / K / L.", "Press ⌃⌘⇧= to equalise the splits."],
         check="You can focus on one pane and then get the grid back."),
    dict(id="m2-3", m="m2", level=1, title="Tabs (surfaces) in a pane",
         steps=["⌘T opens a new surface in the current pane.", "⌘R renames it.", "Switch with ⌃1..9 or ⌘⇧[ and ⌘⇧].", "⌘⇧T reopens the last closed one."],
         check="One pane holds two named tabs and you can flip between them."),
    dict(id="m2-4", m="m2", level=1, title="Workspaces",
         steps=["⌘N creates a workspace; ⌘⇧R renames it.", "Jump with ⌘1..9.", "Reorder with ⌃⌥⌘[ and ⌃⌥⌘]."],
         check="You have a workspace named for a project and can jump to it by number."),
    dict(id="m2-5", m="m2", level=2, title="Move a surface between panes",
         steps=["Make two panes.", "Press ⌥⌘⇧ with an arrow key to push the current surface into the neighbouring pane."],
         check="A tab moved panes without being closed."),
    dict(id="m2-6", m="m2", level=2, title="Group workspaces",
         steps=["Create three workspaces.", "Select two and press ⌘⇧G.", "Collapse the group with ⌃⌘.", "Run `cmux workspace-group list`."],
         check="Two workspaces sit under one collapsible, named group."),
    dict(id="m2-7", m="m2", level=2, title="Build a layout from the CLI",
         steps=["Run `cmux new-workspace --name scratch --cwd ~`.", "Run `cmux new-split right` then `cmux new-split down --command \"ls\"`.", "Run `cmux tree` to see the result.", "Clean up with `cmux close-workspace --workspace <ref>`."],
         check="You made the same kind of layout as m2-1 without touching the keyboard shortcuts."),
    dict(id="m2-8", m="m2", level=2, title="Save a layout template",
         steps=["Arrange a layout you like.", "Press ⌃⌘S to save it as a template.", "Open a new workspace from that template."],
         check="A saved layout comes back on demand."),
    # ---- Module 3
    dict(id="m3-1", m="m3", level=1, title="Get pinged by an agent",
         steps=["Start `claude` in a workspace and give it a task that takes a minute.", "Switch to another workspace.", "When it needs you, press ⌘I to open notifications and ⌘⇧U to jump to the latest unread."],
         check="You were told which workspace needed attention without watching it."),
    dict(id="m3-2", m="m3", level=1, title="Send your own notification",
         steps=["In one workspace run `sleep 5; cmux notify --title \"Done\" --body \"sleep finished\"`.", "Switch away before it fires.", "Then run `cmux list-notifications` and `cmux clear-notifications`."],
         check="You can trigger, read and clear notifications from the shell."),
    dict(id="m3-3", m="m3", level=2, title="Sidebar status pills",
         steps=["Run `cmux set-status build \"passing\" --color \"#2ea043\"`.", "Run `cmux list-status`.", "Run `cmux clear-status build`."],
         check="A coloured pill appeared in the sidebar and then went away."),
    dict(id="m3-4", m="m3", level=2, title="Show progress",
         steps=["Run `for p in 0.2 0.4 0.6 0.8 1.0; do cmux set-progress $p --label \"step $p\"; sleep 1; done; cmux clear-progress`."],
         check="The sidebar progress bar filled up, then cleared."),
    dict(id="m3-5", m="m3", level=2, title="Write to the sidebar log",
         steps=["Run `cmux log --level info --source academy \"hello from the shell\"`.", "Run `cmux list-log --limit 5`."],
         check="Your message shows in the workspace log."),
    dict(id="m3-6", m="m3", level=2, title="Review changes in the diff viewer",
         steps=["In a git repo with uncommitted changes run `cmux diff --unstaged`.", "Move with J and K, search files with /.", "Try `cmux diff --last-turn` after an agent edits files."],
         check="You reviewed a change without leaving cmux."),
    dict(id="m3-7", m="m3", level=2, title="Markdown viewer with live reload",
         steps=["Run `cmux markdown open README.md` (any markdown file).", "Edit the file in another pane and save."],
         check="The formatted viewer updated by itself."),
    dict(id="m3-8", m="m3", level=2, title="Per-workspace todo list",
         steps=["Run `cmux todo add \"read the concepts doc\"`.", "Run `cmux todo list`, then `cmux todo check <id>`.", "In the right sidebar toggle an item with ⌘↩."],
         check="You can keep a checklist attached to a workspace."),
    dict(id="m3-9", m="m3", level=2, title="See how cmux wraps Claude",
         steps=["Run `which claude` and `echo $CMUX_CLAUDE_WRAPPER_SHIM`.", "Run `cmux hooks --help` and read what each agent integration does."],
         check="You understand why notifications work without extra hook scripts."),
    dict(id="m3-10", m="m3", level=3, title="Agent teams",
         steps=["Read https://cmux.com/docs/agent-integrations/claude-code-teams.", "Run `cmux claude-teams` and ask Claude to spin up two teammates on a small task."],
         check="Teammates appear as native cmux splits, with their own sidebar metadata."),
    # ---- Module 4
    dict(id="m4-1", m="m4", level=1, title="Open the browser",
         steps=["Press ⌘⇧L, or run `cmux browser open https://cmux.com/docs`.", "Focus the address bar with ⌘L."],
         check="Docs sit beside your terminal."),
    dict(id="m4-2", m="m4", level=1, title="Browser split",
         steps=["⌥⌘D opens a browser split to the right, ⌥⌘⇧D below.", "Toggle dev tools with ⌥⌘I and the console with ⌥⌘C."],
         check="You can inspect a page without leaving the window."),
    dict(id="m4-3", m="m4", level=1, title="Preview this academy",
         steps=["In the cmux-academy folder run `python3 -m http.server 8000`.", "In another pane run `cmux browser open http://localhost:8000`."],
         check="localhost loads in the in-app browser."),
    dict(id="m4-4", m="m4", level=2, title="Read a page from the shell",
         steps=["Run `cmux browser get title` and `cmux browser url`.", "Run `cmux browser snapshot --interactive` and read the output."],
         check="You can see what a page contains as text."),
    dict(id="m4-5", m="m4", level=2, title="Drive a page",
         steps=["Run `cmux browser eval \"document.title\"`.", "Run `cmux browser wait --text \"cmux\" --timeout-ms 5000`.", "Run `cmux browser screenshot --out ./shot.png`."],
         check="You scripted a page and saved a screenshot."),
    dict(id="m4-6", m="m4", level=2, title="Focus mode and profiles",
         steps=["Press ⌥⌘↩ to enter browser focus mode and again to leave.", "Run `cmux browser profiles list`."],
         check="You know how to give the page your keyboard, and that browser profiles exist."),
    # ---- Module 5
    dict(id="m5-1", m="m5", level=1, title="Use the built-in guides",
         steps=["Run `cmux guide`.", "Run `cmux docs settings` and `cmux docs shortcuts` (they print where to fetch the latest docs)."],
         check="You know the docs are available offline-first from the CLI."),
    dict(id="m5-2", m="m5", level=2, title="Check your config",
         steps=["Run `cmux config path`.", "Back up the file: `cp ~/.config/cmux/cmux.json ~/.config/cmux/cmux.json.bak`.", "Run `cmux config doctor` (or `cmux config validate`)."],
         check="You have a backup and a clean validation result."),
    dict(id="m5-3", m="m5", level=2, title="Change the terminal look",
         steps=["Create `~/.config/ghostty/config` with `font-size = 14` and a `theme = ...` line.", "Run `cmux themes list` to pick a theme.", "Press ⌘⇧, (or run `cmux reload-config`)."],
         check="The change applied with no restart."),
    dict(id="m5-4", m="m5", level=2, title="Rebind a shortcut to a chord",
         steps=["In cmux.json add `\"shortcuts\": { \"bindings\": { \"showNotifications\": [\"ctrl+b\", \"i\"] } }`.", "Reload config.", "Press ⌃B then I."],
         check="A two-step chord opens the notifications panel."),
    dict(id="m5-5", m="m5", level=2, title="Type into another terminal",
         steps=["Make a second pane and find its ref with `cmux tree`.", "Run `cmux send --surface <ref> \"echo hi\"` then `cmux send-key --surface <ref> enter`.", "Run `cmux read-screen --surface <ref> --lines 10`."],
         check="You ran a command in a pane you were not focused on and read the result."),
    dict(id="m5-6", m="m5", level=3, title="Write a custom workspace command",
         steps=["Copy the Review Setup example from https://cmux.com/docs/custom-commands into cmux.json.", "Change the cwd and commands to your project.", "Reload config and run it from the command palette."],
         check="One palette entry opens your whole working layout."),
    dict(id="m5-7", m="m5", level=3, title="Remote workspace over SSH",
         steps=["Read https://cmux.com/docs/ssh.", "If you have a host, run `cmux ssh user@host --name \"dev box\"`."],
         check="A remote session lives in a workspace like any other (skip if you have no host)."),
    dict(id="m5-8", m="m5", level=3, title="Capstone: your daily layout",
         steps=["Design a layout: Claude Code on the left, shell top right, browser bottom right.", "Create it with a custom command or `cmux new-workspace --layout <json>`.", "Add a status pill or progress bar from a script."],
         check="You can open your whole setup with one action."),
]

# (category, keys, action, module)
SHORTCUTS = [
    ("App", "⌘ ,", "Open settings", "m1"),
    ("App", "⌘ ⇧ ,", "Reload configuration", "m5"),
    ("App", "⌘ ⇧ P", "Command palette", "m1"),
    ("App", "⌥ ⌘ F", "Global search", "m1"),
    ("App", "⌘ ⇧ N", "New window", "m1"),
    ("App", "⌘ ⇧ O", "Reopen previous session", "m2"),
    ("Workspaces", "⌘ B", "Toggle left sidebar", "m1"),
    ("Workspaces", "⌘ ⌥ B", "Toggle right sidebar", "m1"),
    ("Workspaces", "⌘ N", "New workspace", "m2"),
    ("Workspaces", "⌥ ⌘ N", "New browser workspace", "m4"),
    ("Workspaces", "⌘ P", "Go to workspace", "m1"),
    ("Workspaces", "⌘ 1…9", "Select workspace 1 to 9", "m2"),
    ("Workspaces", "⌃ ⌘ ]", "Next workspace", "m2"),
    ("Workspaces", "⌃ ⌘ [", "Previous workspace", "m2"),
    ("Workspaces", "⌘ ⇧ R", "Rename workspace", "m2"),
    ("Workspaces", "⌘ ⇧ W", "Close workspace", "m2"),
    ("Workspaces", "⌘ ⇧ G", "Group selected workspaces", "m2"),
    ("Workspaces", "⌃ ⌘ .", "Collapse or expand workspace group", "m2"),
    ("Workspaces", "⌃ ⌘ S", "Save workspace layout as template", "m2"),
    ("Workspaces", "⌘ ;", "Mark workspace as done", "m3"),
    ("Workspaces", "⌘ ↩", "Toggle checklist item", "m3"),
    ("Surfaces", "⌘ T", "New surface (tab)", "m2"),
    ("Surfaces", "⌘ W", "Close tab", "m2"),
    ("Surfaces", "⌘ ⇧ T", "Reopen last closed", "m2"),
    ("Surfaces", "⌘ R", "Rename tab", "m2"),
    ("Surfaces", "⌘ ⇧ ]", "Next surface", "m2"),
    ("Surfaces", "⌘ ⇧ [", "Previous surface", "m2"),
    ("Surfaces", "⌃ 1…9", "Select surface 1 to 9", "m2"),
    ("Surfaces", "⌥ ⌘ ⇧ ← → ↑ ↓", "Move surface to the pane in that direction", "m2"),
    ("Surfaces", "⌘ ⇧ M", "Toggle terminal copy mode", "m2"),
    ("Surfaces", "⌘ ⇧ K", "Clear screen", "m2"),
    ("Split panes", "⌘ D", "Split right", "m2"),
    ("Split panes", "⌘ ⇧ D", "Split down", "m2"),
    ("Split panes", "⌥ ⌘ ← → ↑ ↓", "Focus pane in that direction", "m2"),
    ("Split panes", "⌘ ⇧ ↩", "Toggle pane zoom", "m2"),
    ("Split panes", "⌃ ⇧ H J K L", "Resize pane left, down, up, right", "m2"),
    ("Split panes", "⌃ ⌘ ⇧ =", "Equalise split sizes", "m2"),
    ("Split panes", "⌃ ⌘ =", "Increase font size", "m5"),
    ("Split panes", "⌃ ⌘ -", "Decrease font size", "m5"),
    ("Split panes", "⌃ ⌘ 0", "Reset font size", "m5"),
    ("Browser", "⌘ ⇧ L", "Open browser", "m4"),
    ("Browser", "⌥ ⌘ D", "Split browser right", "m4"),
    ("Browser", "⌥ ⌘ ⇧ D", "Split browser down", "m4"),
    ("Browser", "⌘ L", "Focus address bar", "m4"),
    ("Browser", "⌥ ⌘ I", "Toggle developer tools", "m4"),
    ("Browser", "⌥ ⌘ C", "Show JavaScript console", "m4"),
    ("Browser", "⌥ ⌘ ↩", "Enter browser focus mode", "m4"),
    ("Diff viewer", "⌃ ⌘ ⇧ D", "Open diff viewer", "m3"),
    ("Diff viewer", "J / K", "Scroll down / up one step", "m3"),
    ("Diff viewer", "/", "Search diff files", "m3"),
    ("Find", "⌘ F", "Find in terminal", "m2"),
    ("Find", "⌘ ⇧ F", "Find in directory", "m2"),
    ("Notifications", "⌘ I", "Show notifications", "m3"),
    ("Notifications", "⌘ ⇧ U", "Jump to latest unread", "m3"),
    ("Notifications", "⌥ ⌘ U", "Toggle unread state", "m3"),
    ("Canvas", "⌃ ⌘ C", "Toggle canvas layout", "m2"),
    ("Canvas", "⌃ ⌘ T", "Tidy panes into a grid", "m2"),
]

# (command, what it does, module)
COMMANDS = [
    ("cmux tree", "Show windows, workspaces, panes and surfaces as a tree", "m1"),
    ("cmux identify", "Show the ids of the workspace and surface you are in", "m1"),
    ("cmux list-workspaces", "List workspaces in the current window", "m1"),
    ("cmux welcome", "Open the built-in welcome tour", "m1"),
    ("cmux settings path", "Print the path to cmux.json", "m1"),
    ("cmux new-workspace --name N --cwd DIR", "Create a workspace (optionally --command, --layout JSON)", "m2"),
    ("cmux new-split right|left|up|down", "Split the current pane", "m2"),
    ("cmux new-surface --type terminal|browser", "Add a tab to a pane", "m2"),
    ("cmux focus-pane --pane REF", "Focus a pane", "m2"),
    ("cmux rename-workspace TITLE", "Rename a workspace", "m2"),
    ("cmux close-workspace --workspace REF", "Close a workspace", "m2"),
    ("cmux workspace-group create --name N", "Create a workspace group", "m2"),
    ("cmux notify --title T --body B", "Send a notification", "m3"),
    ("cmux list-notifications", "List notifications", "m3"),
    ("cmux jump-to-unread", "Jump to the latest unread notification", "m3"),
    ("cmux clear-notifications", "Clear notifications", "m3"),
    ("cmux set-status KEY VALUE --color #hex", "Set a sidebar status pill", "m3"),
    ("cmux clear-status KEY", "Remove a status pill", "m3"),
    ("cmux set-progress 0.0-1.0 --label T", "Show a sidebar progress bar", "m3"),
    ("cmux log --level L MESSAGE", "Append to the workspace log", "m3"),
    ("cmux diff --unstaged|--staged|--last-turn", "Open the diff viewer on a git source", "m3"),
    ("cmux markdown open FILE", "Open a markdown file in the live viewer", "m3"),
    ("cmux todo add|list|check TEXT", "Per-workspace checklist", "m3"),
    ("cmux claude-teams", "Launch Claude Code with teammates as native splits", "m3"),
    ("cmux browser open URL", "Open a browser split", "m4"),
    ("cmux browser snapshot --interactive", "Dump the page as text with clickable refs", "m4"),
    ("cmux browser eval SCRIPT", "Run JavaScript in the page", "m4"),
    ("cmux browser click|fill|type SELECTOR", "Interact with the page", "m4"),
    ("cmux browser screenshot --out PATH", "Save a screenshot", "m4"),
    ("cmux send --surface REF TEXT", "Type text into a terminal", "m5"),
    ("cmux send-key --surface REF KEY", "Press a key in a terminal", "m5"),
    ("cmux read-screen --surface REF --lines N", "Read what is on a terminal's screen", "m5"),
    ("cmux config doctor|validate|path", "Check and locate configuration", "m5"),
    ("cmux reload-config", "Reload Ghostty config and cmux.json without a restart", "m5"),
    ("cmux themes list|set", "Browse and set terminal themes", "m5"),
    ("cmux ssh USER@HOST", "Open a remote workspace over SSH", "m5"),
    ("cmux guide", "Print the agent-oriented guide", "m5"),
    ("cmux docs settings|shortcuts|api|browser", "Show how to fetch the latest docs", "m5"),
]

# (title, url, kind, module, why)
LINKS = [
    ("Getting started", "https://cmux.com/docs/getting-started", "Docs", "m1", "Install, first workspace, first CLI commands."),
    ("Concepts", "https://cmux.com/docs/concepts", "Docs", "m1", "The window, workspace, pane, surface, panel hierarchy. Read this first."),
    ("Keyboard shortcuts", "https://cmux.com/docs/keyboard-shortcuts", "Docs", "m2", "The full default list, plus how to rebind and chain shortcuts."),
    ("Workspace groups", "https://cmux.com/docs/workspace-groups", "Docs", "m2", "Collapsible named sections for related workspaces."),
    ("Notifications", "https://cmux.com/docs/notifications", "Docs", "m3", "How agent notifications and the notification panel work."),
    ("Claude Code teams", "https://cmux.com/docs/agent-integrations/claude-code-teams", "Docs", "m3", "Teammate agents as native splits."),
    ("API and CLI reference", "https://cmux.com/docs/api", "Docs", "m3", "Every command, with the socket equivalent."),
    ("Browser automation", "https://cmux.com/docs/browser-automation", "Docs", "m4", "Navigation, snapshots, clicks, console and error access."),
    ("Configuration", "https://cmux.com/docs/configuration", "Docs", "m5", "Config file locations and the cmux.json template."),
    ("Custom commands", "https://cmux.com/docs/custom-commands", "Docs", "m5", "Define workspace layouts and palette actions."),
    ("SSH and remote", "https://cmux.com/docs/ssh", "Docs", "m5", "Remote workspaces over SSH or Mosh."),
    ("cmux TUI", "https://cmux.com/docs/tui", "Docs", "m5", "Session model and the headless control socket."),
    ("cmux on GitHub", "https://github.com/manaflow-ai/cmux", "Source", "m5", "Source, issues, and the cmux.json schema."),
    ("Ghostty configuration", "https://ghostty.org/docs/config", "Docs", "m5", "cmux uses Ghostty config for fonts, themes and terminal behaviour."),
    ("reveal.js documentation", "https://revealjs.com", "Docs", "m1", "How the slide deck in this academy is built, if you want to edit it."),
]


def write_data_js():
    payload = dict(modules=MODULES, tasks=TASKS,
                   shortcuts=[dict(cat=c, keys=k, action=a, m=m) for c, k, a, m in SHORTCUTS],
                   commands=[dict(cmd=c, what=w, m=m) for c, w, m in COMMANDS],
                   links=[dict(title=t, url=u, kind=k, m=m, why=w) for t, u, k, m, w in LINKS])
    out = ROOT / "assets" / "data.js"
    out.write_text("// Generated by tools/build.py. Edit the data there, not here.\nwindow.ACADEMY = "
                   + json.dumps(payload, ensure_ascii=False, indent=1) + ";\n", encoding="utf-8")
    print(f"wrote {out.relative_to(ROOT)}: {len(TASKS)} tasks, {len(SHORTCUTS)} shortcuts, "
          f"{len(COMMANDS)} commands, {len(LINKS)} links")


def write_xlsx():
    try:
        from openpyxl import Workbook
        from openpyxl.formatting.rule import CellIsRule, FormulaRule
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError:
        print("openpyxl not installed: skipped the spreadsheet (pip install openpyxl)")
        return

    mod_title = {m["id"]: f'{m["n"]}. {m["title"]}' for m in MODULES}
    level_name = {1: "Beginner", 2: "Intermediate", 3: "Stretch"}
    head_fill = PatternFill("solid", fgColor="1F2937")
    head_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="D1D5DB")
    border = Border(top=thin, bottom=thin, left=thin, right=thin)
    wrap = Alignment(wrap_text=True, vertical="top")

    wb = Workbook()

    def sheet(title, headers, rows, widths, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        ws.append(headers)
        for r in rows:
            ws.append(list(r))
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=c)
            cell.fill, cell.font = head_fill, head_font
            cell.alignment = Alignment(vertical="center")
            ws.column_dimensions[get_column_letter(c)].width = widths[c - 1]
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(headers)):
            for cell in row:
                cell.alignment, cell.border = wrap, border
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"
        return ws

    # --- Start here
    ws = wb.active
    ws.title = "Start here"
    lines = [
        ("cmux Academy tracker", True),
        ("", False),
        ("How to use this workbook", True),
        ("1. Work through the modules in order (see the Progress tab). The site has the lessons: open index.html.", False),
        ("2. Tasks tab: set Status to Doing or Done as you go. Progress updates by itself.", False),
        ("3. Shortcuts tab: rate your confidence 1 (never used) to 5 (muscle memory). Sort by confidence and drill the low ones.", False),
        ("4. CLI tab: mark each command Tried once you have run it.", False),
        ("5. Reading list: mark pages Read, and put what surprised you in Notes.", False),
        ("", False),
        ("Rule of thumb: a shortcut is learned when you use it without thinking for three days running.", False),
        ("The checklist on the website and this workbook are separate trackers; use whichever suits you.", False),
    ]
    for i, (text, bold) in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=text)
        c.font = Font(bold=bold, size=14 if i == 1 else 11)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 110

    # --- Tasks
    task_rows = [(t["id"], mod_title[t["m"]], t["title"], level_name[t["level"]],
                  " / ".join(t["steps"]), t["check"], "Todo", None, None) for t in TASKS]
    tws = sheet("Tasks", ["ID", "Module", "Task", "Level", "Steps", "You know it worked when", "Status", "Date done", "Notes"],
                task_rows, [8, 24, 30, 13, 70, 44, 11, 12, 36])
    n_tasks = len(TASKS)
    dv = DataValidation(type="list", formula1='"Todo,Doing,Done"', allow_blank=False)
    tws.add_data_validation(dv)
    dv.add(f"G2:G{n_tasks + 1}")
    for col in "H":
        for r in range(2, n_tasks + 2):
            tws[f"{col}{r}"].number_format = "yyyy-mm-dd"
    rng = f"A2:I{n_tasks + 1}"
    tws.conditional_formatting.add(rng, FormulaRule(formula=['$G2="Done"'], fill=PatternFill("solid", bgColor="DCFCE7")))
    tws.conditional_formatting.add(rng, FormulaRule(formula=['$G2="Doing"'], fill=PatternFill("solid", bgColor="FEF9C3")))

    # --- Progress (formulas over Tasks)
    prog_rows = []
    for i, m in enumerate(MODULES, start=2):
        label = mod_title[m["id"]]
        prog_rows.append((label, m["minutes"],
                          f'=COUNTIF(Tasks!$B$2:$B${n_tasks + 1},A{i})',
                          f'=COUNTIFS(Tasks!$B$2:$B${n_tasks + 1},A{i},Tasks!$G$2:$G${n_tasks + 1},"Done")',
                          f'=IF(C{i}=0,0,D{i}/C{i})'))
    last = len(MODULES) + 1
    prog_rows.append(("Total", f"=SUM(B2:B{last})", f"=SUM(C2:C{last})", f"=SUM(D2:D{last})",
                      f"=IF(C{last + 1}=0,0,D{last + 1}/C{last + 1})"))
    pws = sheet("Progress", ["Module", "Est. minutes", "Tasks", "Done", "% complete"], prog_rows, [30, 14, 10, 10, 14])
    for r in range(2, last + 2):
        pws[f"E{r}"].number_format = "0%"
    for c in range(1, 6):
        pws.cell(row=last + 1, column=c).font = Font(bold=True)
    pws.auto_filter.ref = None
    pws.conditional_formatting.add(f"E2:E{last + 1}", CellIsRule(operator="greaterThanOrEqual", formula=["1"],
                                   fill=PatternFill("solid", bgColor="DCFCE7")))
    pws.sheet_properties.tabColor = "2EA043"

    # --- Shortcuts
    sc_rows = [(c, k, a, mod_title[m], None, None) for c, k, a, m in SHORTCUTS]
    sws = sheet("Shortcuts", ["Category", "Keys", "Action", "Module", "Confidence (1-5)", "Last practised"],
                sc_rows, [16, 22, 44, 24, 17, 15])
    dv2 = DataValidation(type="list", formula1='"1,2,3,4,5"', allow_blank=True)
    sws.add_data_validation(dv2)
    dv2.add(f"E2:E{len(SHORTCUTS) + 1}")
    for r in range(2, len(SHORTCUTS) + 2):
        sws[f"F{r}"].number_format = "yyyy-mm-dd"
        sws[f"B{r}"].font = Font(name="Menlo", bold=True)
    sws.conditional_formatting.add(f"E2:E{len(SHORTCUTS) + 1}", CellIsRule(operator="lessThanOrEqual", formula=["2"],
                                   fill=PatternFill("solid", bgColor="FEE2E2")))
    sws.conditional_formatting.add(f"E2:E{len(SHORTCUTS) + 1}", CellIsRule(operator="greaterThanOrEqual", formula=["4"],
                                   fill=PatternFill("solid", bgColor="DCFCE7")))

    # --- CLI
    cl_rows = [(c, w, mod_title[m], "No", None) for c, w, m in COMMANDS]
    cws = sheet("CLI", ["Command", "What it does", "Module", "Tried?", "Notes"], cl_rows, [48, 56, 24, 9, 36])
    dv3 = DataValidation(type="list", formula1='"No,Tried,Comfortable"', allow_blank=False)
    cws.add_data_validation(dv3)
    dv3.add(f"D2:D{len(COMMANDS) + 1}")
    for r in range(2, len(COMMANDS) + 2):
        cws[f"A{r}"].font = Font(name="Menlo")
    cws.conditional_formatting.add(f"D2:D{len(COMMANDS) + 1}", CellIsRule(operator="equal", formula=['"Comfortable"'],
                                   fill=PatternFill("solid", bgColor="DCFCE7")))
    cws.conditional_formatting.add(f"D2:D{len(COMMANDS) + 1}", CellIsRule(operator="equal", formula=['"Tried"'],
                                   fill=PatternFill("solid", bgColor="FEF9C3")))

    # --- Reading list
    rl_rows = [(t, u, k, mod_title[m], w, "To read", None) for t, u, k, m, w in LINKS]
    rws = sheet("Reading list", ["Title", "Link", "Type", "Module", "Why read it", "Status", "Notes"],
                rl_rows, [28, 52, 9, 24, 56, 10, 40])
    dv4 = DataValidation(type="list", formula1='"To read,Reading,Read"', allow_blank=False)
    rws.add_data_validation(dv4)
    dv4.add(f"F2:F{len(LINKS) + 1}")
    for r in range(2, len(LINKS) + 2):
        rws[f"B{r}"].hyperlink = rws[f"B{r}"].value
        rws[f"B{r}"].font = Font(color="1D4ED8", underline="single")

    order = ["Start here", "Progress", "Tasks", "Shortcuts", "CLI", "Reading list"]
    wb._sheets = [wb[name] for name in order]
    out = ROOT / "cmux-academy.xlsx"
    wb.save(out)
    print(f"wrote {out.relative_to(ROOT)}: sheets {wb.sheetnames}")


if __name__ == "__main__":
    write_data_js()
    write_xlsx()
    sys.exit(0)
